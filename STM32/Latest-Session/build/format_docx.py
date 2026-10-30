import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

INTERMEDIATE_DOCX = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "intermediate.docx"))
FINAL_DOCX = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "STM32_Complete_Professional_Reference.docx"))

# Color Palette
COLOR_NAVY_PRIMARY   = RGBColor(0x0F, 0x25, 0x37)  # #0F2537
COLOR_BLUE_SECONDARY = RGBColor(0x1E, 0x4E, 0x79)  # #1E4E79
COLOR_TEAL_ACCENT    = RGBColor(0x00, 0x80, 0x80)  # #008080
COLOR_DARK_TEXT      = RGBColor(0x22, 0x22, 0x22)  # #222222
COLOR_MUTED_GRAY     = RGBColor(0x66, 0x66, 0x66)  # #666666

HEX_NAVY_PRIMARY     = "0F2537"
HEX_BLUE_SECONDARY   = "1E4E79"
HEX_TEAL_ACCENT      = "008080"
HEX_LIGHT_BG         = "F8FAFC"
HEX_BORDER_GRAY      = "D0D7DE"
HEX_CODE_BG          = "F6F8FA"

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>'
        % (nsdecls('w'), top, bottom, left, right)
    )
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd %s w:fill="%s" w:val="clear"/>' % (nsdecls('w'), color_hex))
    tcPr.append(shd)

def set_cell_left_border_only(cell, color_hex, size_eighth_pt=28): # 28 = 3.5 pt
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        r'<w:tcBorders %s><w:top w:val="none"/><w:left w:val="single" w:sz="%d" w:space="0" w:color="%s"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>'
        % (nsdecls('w'), size_eighth_pt, color_hex)
    )
    tcPr.append(borders)

def set_table_borders(table, color_hex="D0D7DE"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders %s><w:top w:val="single" w:sz="6" w:space="0" w:color="%s"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="%s"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/><w:insideV w:val="none"/></w:tblBorders>'
        % (nsdecls('w'), color_hex, HEX_NAVY_PRIMARY, color_hex)
    )
    tblPr.append(borders)

def add_header_footer(doc):
    for s_idx, section in enumerate(doc.sections):
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.different_first_page_header_footer = True
        
        # Header (Pages 2+)
        header = section.header
        hp = header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("STM32 Complete Professional Reference  |  Consolidated Technical Reference")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = COLOR_MUTED_GRAY
        
        # Add bottom border to header
        hpPr = hp._p.get_or_add_pPr()
        hpBrd = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="4" w:space="4" w:color="%s"/></w:pBdr>' % (nsdecls('w'), HEX_TEAL_ACCENT))
        hpPr.append(hpBrd)
        
        # Footer (Pages 2+)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = ""
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        # Left side: Organization
        frun_left = fp.add_run("Assiut Robotics Embedded Systems Batch\t")
        frun_left.font.name = "Segoe UI"
        frun_left.font.size = Pt(8.5)
        frun_left.font.color.rgb = COLOR_MUTED_GRAY
        
        # Right side: Page X of Y
        frun_p = fp.add_run("Page ")
        frun_p.font.name = "Segoe UI"
        frun_p.font.size = Pt(8.5)
        frun_p.font.color.rgb = COLOR_MUTED_GRAY
        
        # Insert PAGE field
        fld1 = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fld1)
        
        frun_of = fp.add_run(" of ")
        frun_of.font.name = "Segoe UI"
        frun_of.font.size = Pt(8.5)
        frun_of.font.color.rgb = COLOR_MUTED_GRAY
        
        # Insert NUMPAGES field
        fld2 = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
        fp._p.append(fld2)

def build_cover_page(doc):
    # Insert at top of document
    p_top = doc.paragraphs[0]
    
    # 1. Dark Navy Accent Block Table for Cover Title
    cover_table = doc.add_table(rows=1, cols=1)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_cell = cover_table.rows[0].cells[0]
    set_cell_shading(cover_cell, HEX_NAVY_PRIMARY)
    set_cell_margins(cover_cell, top=360, bottom=360, left=300, right=300)
    
    # Teal left accent border on cover box
    set_cell_left_border_only(cover_cell, HEX_TEAL_ACCENT, size_eighth_pt=48) # 6 pt thick
    
    cp = cover_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(8)
    
    c_badge = cp.add_run("ARM CORTEX-M EMBEDDED SYSTEMS TECHNICAL TRAINING\n")
    c_badge.font.name = "Segoe UI"
    c_badge.font.size = Pt(11)
    c_badge.font.bold = True
    c_badge.font.color.rgb = RGBColor(0x00, 0xD4, 0xD4) # Bright Cyan accent
    
    c_title = cp.add_run("STM32 Complete Professional Reference\n")
    c_title.font.name = "Segoe UI"
    c_title.font.size = Pt(26)
    c_title.font.bold = True
    c_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    c_sub = cp.add_run("The Definitive Technical Textbook & Practical Training Manual")
    c_sub.font.name = "Segoe UI"
    c_sub.font.size = Pt(13)
    c_sub.font.color.rgb = RGBColor(0xDD, 0xEE, 0xFF)
    
    # Move cover table to top
    p_top._p.addprevious(cover_table._tbl)
    
    # 2. Metadata / Summary box below
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(24)
    meta_p.paragraph_format.space_after = Pt(18)
    meta_p.paragraph_format.line_spacing = 1.2
    
    m_run = meta_p.add_run(
        "A rigorous, hardware-level and HAL-level consolidation of STM32 microcontrollers, development workflows, "
        "direct memory access (DMA), vectored interrupts, multi-channel timers, serial communication protocols (UART/SPI/I²C), "
        "analog-to-digital conversion, reusable modular driver architecture, digital signal processing, and FreeRTOS multitasking.\n"
    )
    m_run.font.name = "Segoe UI"
    m_run.font.size = Pt(11)
    m_run.font.color.rgb = COLOR_DARK_TEXT
    
    p_top._p.addprevious(meta_p._p)
    
    # 3. Specification Table
    info_table = doc.add_table(rows=5, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_table, HEX_BORDER_GRAY)
    
    meta_data = [
        ("Institutional Batch", "Assiut Robotics Embedded Systems Training Program"),
        ("Baseline Toolchain", "STM32CubeIDE 2.2.0 • STM32CubeMX 6.18.1 • GNU ARM Toolchain"),
        ("Primary Target Silicon", "STM32F103C8T6 (Cortex-M3, 72 MHz) with STM32F4/G4/H7 extensions"),
        ("Document Edition", "Version 5.0 (Consolidated Master Edition)"),
        ("Status", "Production Technical Reference — Verified Against Official Datasheets & RMs")
    ]
    
    for row_idx, (k, v) in enumerate(meta_data):
        row = info_table.rows[row_idx]
        c0, c1 = row.cells[0], row.cells[1]
        set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
        set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
        set_cell_shading(c0, HEX_LIGHT_BG)
        set_cell_shading(c1, "FFFFFF")
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(1)
        p0.paragraph_format.space_after = Pt(1)
        r0 = p0.add_run(k)
        r0.font.name = "Segoe UI"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_NAVY_PRIMARY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(v)
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK_TEXT
        
    p_top._p.addprevious(info_table._tbl)
    
    # 4. Page Break after cover page
    pb_p = doc.add_paragraph()
    pb_p.paragraph_format.page_break_before = True
    p_top._p.addprevious(pb_p._p)

def format_headings_and_body(doc):
    # Regex to detect major numbered chapters for page breaks
    major_chapter_pattern = re.compile(r'^(#|\d+\.\s+|PART\s+\d+|Appendix\s+[A-Z])', re.IGNORECASE)
    
    for p in doc.paragraphs:
        style_name = p.style.name
        text = p.text.strip()
        
        if style_name == "Heading 1":
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            
            # If it's a major chapter or appendix, start on fresh page!
            if (re.match(r'^\d+\.', text) or text.startswith("Appendix")) and not text.startswith("1.1"):
                p.paragraph_format.page_break_before = True
                
            for run in p.runs:
                run.font.name = "Segoe UI"
                run.font.size = Pt(18)
                run.font.bold = True
                run.font.color.rgb = COLOR_NAVY_PRIMARY
                
            # Add subtle teal bottom border to Heading 1
            pPr = p._p.get_or_add_pPr()
            pBrd = parse_xml(r'<w:pBdr %s><w:bottom w:val="single" w:sz="8" w:space="4" w:color="%s"/></w:pBdr>' % (nsdecls('w'), HEX_TEAL_ACCENT))
            pPr.append(pBrd)

        elif style_name == "Heading 2":
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.name = "Segoe UI"
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.color.rgb = COLOR_BLUE_SECONDARY

        elif style_name == "Heading 3":
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.name = "Segoe UI"
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = COLOR_TEAL_ACCENT

        elif style_name == "Source Code":
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.left_indent = Inches(0.25)
            # Background shading for code blocks
            pPr = p._p.get_or_add_pPr()
            shd = parse_xml(r'<w:shd %s w:fill="%s" w:val="clear"/>' % (nsdecls('w'), HEX_CODE_BG))
            pPr.append(shd)
            # Left border on code
            pBrd = parse_xml(r'<w:pBdr %s><w:left w:val="single" w:sz="12" w:space="8" w:color="%s"/></w:pBdr>' % (nsdecls('w'), HEX_TEAL_ACCENT))
            pPr.append(pBrd)
            
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

        elif style_name in ["Body Text", "Normal", "First Paragraph", "Compact"]:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            for run in p.runs:
                if not run.font.name:
                    run.font.name = "Segoe UI"
                if not run.font.size:
                    run.font.size = Pt(10)
                if not run.font.color.rgb:
                    run.font.color.rgb = COLOR_DARK_TEXT
                    
        elif style_name == "Captioned Figure" or style_name == "Image Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                run.font.name = "Segoe UI"
                run.font.size = Pt(9)
                run.font.italic = True
                run.font.color.rgb = COLOR_MUTED_GRAY

def format_all_tables(doc):
    print("Styling tables...")
    for t_idx, table in enumerate(doc.tables):
        # Center table
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, HEX_BORDER_GRAY)
        
        # Mark header row
        header_tr = table.rows[0]._tr
        trPr = header_tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
        
        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
            
            is_header = (r_idx == 0)
            bg_color = HEX_NAVY_PRIMARY if is_header else (HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF")
            
            for cell in row.cells:
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                set_cell_shading(cell, bg_color)
                
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.05
                    for run in p.runs:
                        run.font.name = "Segoe UI"
                        if is_header:
                            run.font.size = Pt(9.5)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            run.font.size = Pt(9)
                            run.font.color.rgb = COLOR_DARK_TEXT

def transform_callout_boxes(doc):
    print("Formatting callouts and alerts...")
    # Map alerts to colors and headers
    alert_configs = {
        "IMPORTANT": (HEX_NAVY_PRIMARY, "F0F4F8", "KEY CONCEPT: ", RGBColor(0x0F, 0x25, 0x37)),
        "NOTE":      (HEX_BLUE_SECONDARY, "F0F4F8", "ENGINEERING NOTE: ", RGBColor(0x1E, 0x4E, 0x79)),
        "TIP":       (HEX_TEAL_ACCENT, "F0FAF8", "PRACTICAL BENCH TIP: ", RGBColor(0x00, 0x80, 0x80)),
        "WARNING":   ("D97706", "FFFBEB", "WARNING: ", RGBColor(0xD9, 0x77, 0x06)),
        "CAUTION":   ("B91C1C", "FEF2F2", "COMMON MISTAKE / CAUTION: ", RGBColor(0xB9, 0x1C, 0x1C))
    }
    
    for p in doc.paragraphs:
        text = p.text.strip()
        m = re.match(r'^\[!(IMPORTANT|NOTE|TIP|WARNING|CAUTION)\]\s*(.*)$', text, re.DOTALL)
        if m:
            alert_type = m.group(1)
            content = m.group(2).strip()
            border_hex, bg_hex, label_prefix, label_color = alert_configs[alert_type]
            
            # Format paragraph directly into a sleek callout block
            pPr = p._p.get_or_add_pPr()
            # Background shading
            shd = parse_xml(r'<w:shd %s w:fill="%s" w:val="clear"/>' % (nsdecls('w'), bg_hex))
            pPr.append(shd)
            # Left border 28 = 3.5 pt
            pBrd = parse_xml(r'<w:pBdr %s><w:left w:val="single" w:sz="28" w:space="12" w:color="%s"/></w:pBdr>' % (nsdecls('w'), border_hex))
            pPr.append(pBrd)
            
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            
            # Rewrite runs
            p.text = ""
            r_hdr = p.add_run(label_prefix)
            r_hdr.font.name = "Segoe UI"
            r_hdr.font.size = Pt(10)
            r_hdr.font.bold = True
            r_hdr.font.color.rgb = label_color
            
            r_body = p.add_run(content)
            r_body.font.name = "Segoe UI"
            r_body.font.size = Pt(9.5)
            r_body.font.color.rgb = COLOR_DARK_TEXT

def main():
    print(f"Loading {INTERMEDIATE_DOCX}...")
    doc = docx.Document(INTERMEDIATE_DOCX)
    
    print("Applying page setup and headers/footers...")
    add_header_footer(doc)
    
    print("Formatting headings and typography...")
    format_headings_and_body(doc)
    
    print("Styling tables...")
    format_all_tables(doc)
    
    print("Transforming callouts...")
    transform_callout_boxes(doc)
    
    print("Building executive cover page...")
    build_cover_page(doc)
    
    print(f"Saving final professional DOCX to {FINAL_DOCX}...")
    doc.save(FINAL_DOCX)
    print("Final DOCX built successfully!")

if __name__ == "__main__":
    main()
